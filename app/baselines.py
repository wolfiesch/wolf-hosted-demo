"""The same replace-text edit, written the way an agent usually writes it with the common
Python libraries (openpyxl, python-docx, python-pptx): load, replace inside each run or cell,
save.

    pip install openpyxl==3.1.5 python-docx==1.2.0 python-pptx==1.0.2
    python3 baselines.py <input> <output> <find> <replace>

Prints one JSON line: {"library": "<name> <version>", "replacements": <n>}.
"""

from __future__ import annotations

import json
import sys
from importlib.metadata import version
from pathlib import Path


def _replace_runs(paragraphs, find: str, replace: str) -> int:
    count = 0
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if find in run.text:
                count += run.text.count(find)
                run.text = run.text.replace(find, replace)
    return count


def edit_docx(source: Path, target: Path, find: str, replace: str) -> dict:
    import docx

    document = docx.Document(str(source))
    count = _replace_runs(document.paragraphs, find, replace)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                count += _replace_runs(cell.paragraphs, find, replace)
    document.save(str(target))
    return {"library": f"python-docx {version('python-docx')}", "replacements": count}


def edit_pptx(source: Path, target: Path, find: str, replace: str) -> dict:
    import pptx

    presentation = pptx.Presentation(str(source))
    count = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                count += _replace_runs(shape.text_frame.paragraphs, find, replace)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        count += _replace_runs(cell.text_frame.paragraphs, find, replace)
    presentation.save(str(target))
    return {"library": f"python-pptx {version('python-pptx')}", "replacements": count}


def edit_xlsx(source: Path, target: Path, find: str, replace: str) -> dict:
    import openpyxl

    workbook = openpyxl.load_workbook(str(source))
    count = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and find in cell.value:
                    count += cell.value.count(find)
                    cell.value = cell.value.replace(find, replace)
    workbook.save(str(target))
    return {"library": f"openpyxl {version('openpyxl')}", "replacements": count}


EDITORS = {"docx": edit_docx, "pptx": edit_pptx, "xlsx": edit_xlsx}


def main() -> int:
    source, target, find, replace = sys.argv[1:5]
    source_path, target_path = Path(source), Path(target)
    result = EDITORS[source_path.suffix.lstrip(".").lower()](
        source_path, target_path, find, replace
    )
    print(json.dumps(result))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
